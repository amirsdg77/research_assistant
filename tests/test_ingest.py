"""Tests for src.ingest and the search_documents tool.

PDFs are tricky to generate without an additional dependency. We test the
PDF code path by monkeypatching `pypdf.PdfReader` to return a synthetic
page list — that exercises our enumeration logic without requiring a real
text-bearing PDF.

DOCX is tested with python-docx writing a real file in-memory.
TXT/MD use plain bytes.

The full ingest pipeline is tested with mocked memory and DB so this file
doesn't need Chroma or Postgres running.
"""
from __future__ import annotations

import io
from contextlib import asynccontextmanager
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest

import src.ingest as ingest_mod
from src import memory as memory_mod
from src.ingest import (
    IngestError,
    ParsedPage,
    chunks_from_pages,
    detect_content_type,
    ingest_document,
    parse,
)


# --- content-type detection -------------------------------------------------


def test_detect_content_type_by_extension():
    assert detect_content_type("a.pdf") == "application/pdf"
    assert detect_content_type("b.txt") == "text/plain"
    assert detect_content_type("c.md") == "text/markdown"
    assert (
        detect_content_type("d.docx")
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_detect_content_type_unknown_uses_fallback():
    assert detect_content_type("weird.bin", fallback="text/plain") == "text/plain"


def test_detect_content_type_unknown_no_fallback_raises():
    with pytest.raises(IngestError):
        detect_content_type("weird.bin")


# --- parsers ---------------------------------------------------------------


def test_parse_txt_returns_single_page():
    pages = parse(b"Hello, plain text.", "text/plain")
    assert pages == [ParsedPage(text="Hello, plain text.", page=None)]


def test_parse_md_returns_single_page():
    pages = parse(b"# Heading\n\nBody.", "text/markdown")
    assert len(pages) == 1
    assert "Heading" in pages[0].text


def test_parse_txt_empty_returns_no_pages():
    assert parse(b"   ", "text/plain") == []


def test_parse_docx_extracts_paragraphs():
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    doc.save(buf)

    pages = parse(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert len(pages) == 1  # DOCX concatenates into one block
    assert "First paragraph" in pages[0].text
    assert "Second paragraph" in pages[0].text
    assert pages[0].page is None


def test_parse_pdf_reads_pages(monkeypatch):
    """We patch pypdf.PdfReader to return a fake page list; the goal is to
    verify our enumeration + page-number logic, not pypdf itself."""

    class _FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class _FakeReader:
        def __init__(self, _stream):
            self.pages = [_FakePage("Page one body."), _FakePage("Page two body.")]

    monkeypatch.setattr("pypdf.PdfReader", _FakeReader)
    pages = parse(b"%PDF-fake", "application/pdf")
    assert len(pages) == 2
    assert pages[0].page == 1
    assert pages[1].page == 2
    assert "Page one" in pages[0].text


def test_parse_pdf_skips_blank_pages(monkeypatch):
    class _FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class _FakeReader:
        def __init__(self, _stream):
            self.pages = [_FakePage("Body."), _FakePage(""), _FakePage("More body.")]

    monkeypatch.setattr("pypdf.PdfReader", _FakeReader)
    pages = parse(b"%PDF-fake", "application/pdf")
    # Blanks dropped; surviving pages keep their original page numbers.
    assert len(pages) == 2
    assert pages[0].page == 1
    assert pages[1].page == 3


def test_parse_pdf_handles_broken_input(monkeypatch):
    class _FakeReader:
        def __init__(self, _stream):
            raise ValueError("not a PDF")

    monkeypatch.setattr("pypdf.PdfReader", _FakeReader)
    with pytest.raises(IngestError):
        parse(b"garbage", "application/pdf")


def test_parse_unsupported_content_type_raises():
    with pytest.raises(IngestError):
        parse(b"x", "application/zip")


# --- chunks_from_pages ----------------------------------------------------


def test_chunks_from_pages_preserves_page_metadata():
    pages = [
        ParsedPage(text="alpha", page=1),
        ParsedPage(text="beta", page=2),
    ]
    specs = chunks_from_pages(pages)
    assert len(specs) == 2
    assert specs[0].page == 1 and specs[0].chunk_index == 0
    assert specs[1].page == 2 and specs[1].chunk_index == 1


def test_chunks_from_pages_handles_missing_page():
    pages = [ParsedPage(text="alpha"), ParsedPage(text="beta")]
    specs = chunks_from_pages(pages)
    assert all(s.page is None for s in specs)


def test_chunks_from_pages_empty_input():
    assert chunks_from_pages([]) == []


def test_chunks_from_pages_skips_pages_that_chunk_to_empty():
    """A page with only whitespace produces no chunks; indices stay continuous
    across the gap."""
    pages = [ParsedPage(text="alpha", page=1), ParsedPage(text="   ", page=2),
             ParsedPage(text="beta", page=3)]
    specs = chunks_from_pages(pages)
    assert len(specs) == 2
    assert [s.chunk_index for s in specs] == [0, 1]
    assert [s.page for s in specs] == [1, 3]


# --- ingest_document end-to-end -------------------------------------------


@pytest.fixture
def fake_chroma_store(monkeypatch):
    """Install a fake ChromaMemoryStore that records add_document_chunks calls."""
    from src.memory import ChromaMemoryStore

    store = ChromaMemoryStore()
    store._client = SimpleNamespace()  # sentinel so isinstance check passes
    store.add_document_chunks = AsyncMock(return_value=3)  # type: ignore[method-assign]
    monkeypatch.setattr(memory_mod, "store", store)
    return store


@pytest.fixture
def stub_db_session(monkeypatch):
    """Replace session_scope with a recording stub: tracks objects added and
    pretends quota queries return an empty list."""
    added: list = []

    class _StubSession:
        def add(self, obj):
            added.append(obj)

        async def execute(self, _stmt):
            # Returns a fake Result with `.scalars().all()` → [] (no existing docs).
            return SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: []))

        async def commit(self):
            ...

        async def rollback(self):
            ...

    @asynccontextmanager
    async def _scope():
        yield _StubSession()

    monkeypatch.setattr(ingest_mod, "session_scope", _scope)
    return added


async def test_ingest_txt_end_to_end(fake_chroma_store, stub_db_session):
    sid = uuid4()
    result = await ingest_document(
        filename="note.txt",
        data=b"This is the content.",
        session_id=sid,
    )
    assert result.chunk_count == 3
    # Memory store received the chunks.
    fake_chroma_store.add_document_chunks.assert_awaited_once()
    kwargs = fake_chroma_store.add_document_chunks.await_args.kwargs
    assert kwargs["session_id"] == sid
    assert kwargs["filename"] == "note.txt"
    assert len(list(kwargs["chunks"])) >= 1
    # A Document row was added with the right shape.
    assert len(stub_db_session) == 1
    doc = stub_db_session[0]
    assert doc.filename == "note.txt"
    assert doc.session_id == sid
    assert doc.chunk_count == 3
    assert doc.byte_size == len(b"This is the content.")


async def test_ingest_global_document_uses_none_session(fake_chroma_store, stub_db_session):
    result = await ingest_document(
        filename="g.md",
        data=b"# Global doc",
        session_id=None,
    )
    assert result.chunk_count == 3
    kwargs = fake_chroma_store.add_document_chunks.await_args.kwargs
    assert kwargs["session_id"] is None
    assert stub_db_session[0].session_id is None


async def test_ingest_rejects_oversized_file(fake_chroma_store, stub_db_session, monkeypatch):
    monkeypatch.setattr(ingest_mod.settings, "max_upload_bytes", 10)
    with pytest.raises(IngestError):
        await ingest_document(
            filename="big.txt",
            data=b"x" * 1024,
            session_id=uuid4(),
        )


async def test_ingest_rejects_unsupported_type(fake_chroma_store, stub_db_session):
    with pytest.raises(IngestError):
        await ingest_document(
            filename="archive.zip",
            data=b"PK...",
            session_id=uuid4(),
        )


async def test_ingest_rejects_empty_parse_result(fake_chroma_store, stub_db_session):
    with pytest.raises(IngestError):
        await ingest_document(
            filename="empty.txt",
            data=b"   \n\t",
            session_id=uuid4(),
        )


async def test_ingest_quota_exhausted_raises(monkeypatch, fake_chroma_store):
    """When MAX_DOCS_PER_SESSION is already met, ingestion is refused."""
    sid = uuid4()
    existing_docs = [object()] * 10  # 10 existing >= the default cap

    class _StubSession:
        def add(self, obj):
            ...

        async def execute(self, _stmt):
            return SimpleNamespace(
                scalars=lambda: SimpleNamespace(all=lambda: existing_docs)
            )

        async def commit(self):
            ...

        async def rollback(self):
            ...

    @asynccontextmanager
    async def _scope():
        yield _StubSession()

    monkeypatch.setattr(ingest_mod, "session_scope", _scope)

    with pytest.raises(IngestError):
        await ingest_document(filename="x.txt", data=b"content", session_id=sid)


async def test_ingest_refuses_when_store_is_stub(monkeypatch, stub_db_session):
    """If the Chroma store wasn't installed at startup, ingest must refuse
    rather than silently writing to the stub (which doesn't search documents)."""
    from src.memory import InMemoryStubStore

    monkeypatch.setattr(memory_mod, "store", InMemoryStubStore())
    with pytest.raises(IngestError):
        await ingest_document(
            filename="x.txt", data=b"content", session_id=uuid4()
        )


# --- search_documents tool ------------------------------------------------


async def test_search_documents_without_session_raises():
    from src.tools import invoke
    from src.tools.base import ToolError
    from src.tools.fetch_url import current_session_id

    assert current_session_id.get() is None
    with pytest.raises(ToolError):
        await invoke("search_documents", {"query": "anything"})


async def test_search_documents_returns_hits_with_metadata(monkeypatch):
    from src.memory import MemoryChunk
    from src.tools import invoke
    from src.tools.fetch_url import current_session_id

    async def _fake_search(*, session_id, query, k=5):
        return [
            MemoryChunk(
                chunk="The introduction states X.",
                score=0.9,
                metadata={
                    "filename": "paper.pdf",
                    "page": 2,
                    "document_id": "abc",
                    "source_url": "paper.pdf",
                },
            )
        ]

    fake_store = SimpleNamespace(search_documents=_fake_search)
    monkeypatch.setattr(memory_mod, "store", fake_store)

    sid = uuid4()
    token = current_session_id.set(sid)
    try:
        out = await invoke("search_documents", {"query": "introduction", "k": 3})
    finally:
        current_session_id.reset(token)

    assert len(out.hits) == 1
    hit = out.hits[0]
    assert hit.filename == "paper.pdf"
    assert hit.page == 2
    assert hit.document_id == "abc"
    assert hit.score == 0.9


async def test_search_documents_handles_missing_page_metadata(monkeypatch):
    from src.memory import MemoryChunk
    from src.tools import invoke
    from src.tools.fetch_url import current_session_id

    async def _fake_search(*, session_id, query, k=5):
        return [
            MemoryChunk(
                chunk="text",
                score=0.5,
                metadata={"filename": "note.md", "document_id": "z"},
            )
        ]

    monkeypatch.setattr(memory_mod, "store", SimpleNamespace(search_documents=_fake_search))

    sid = uuid4()
    token = current_session_id.set(sid)
    try:
        out = await invoke("search_documents", {"query": "x"})
    finally:
        current_session_id.reset(token)

    assert out.hits[0].page is None
